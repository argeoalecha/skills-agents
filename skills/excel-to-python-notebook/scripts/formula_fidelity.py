"""Carry a workbook's formulas into its Python conversion, verbatim and executable.

Three Excel-compatible libraries do the work, each for the part it is reliable at:

    openpyxl   reads every formula string exactly as stored, tokenizes it, and writes
               a workbook back out with the formulas still live
    formulas   evaluates those same formula strings with Excel semantics (operator
               precedence, tail conventions, rounding), so the conversion can be
               checked against the formula itself and not only against one cached number
    Excel      remains the authority: wherever its cached value and the engine
               disagree, the cached value wins and the disagreement is reported

Usage:
    python3 formula_fidelity.py manifest <workbook> [--out FILE.json] [--skeleton]
    python3 formula_fidelity.py engine   <workbook>
    python3 formula_fidelity.py export   <workbook> <out.xlsx>

Copy this file next to validate_conversion.py; the Validator imports it.
"""

import argparse
import contextlib
import datetime
import hashlib
import io
import json
import math
import re
import sys
import warnings
import zipfile
from pathlib import Path
from xml.etree import ElementTree

warnings.filterwarnings("ignore")

try:
    import openpyxl
    from openpyxl.formula import Tokenizer
    from openpyxl.formula.tokenizer import Token
    from openpyxl.utils import column_index_from_string
    from openpyxl.utils.cell import coordinate_from_string
    from openpyxl.utils.datetime import to_excel
except ImportError:
    sys.exit("openpyxl required: pip install openpyxl")

# Recalculated on every edit, so a cached value is one frozen draw, not a target.
VOLATILE = {"RAND", "RANDBETWEEN", "RANDARRAY", "NOW", "TODAY", "OFFSET", "INDIRECT", "INFO", "CELL"}

ERROR_VALUES = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NULL!", "#NUM!")

CELL_REF = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
COL_REF = re.compile(r"^(\$?)([A-Za-z]{1,3})$")
ROW_REF = re.compile(r"^(\$?)(\d+)$")
ENGINE_KEY = re.compile(r"^'\[(.+?)\](.+)'!\$?([A-Z]+)\$?([0-9]+)$")


class EngineUnavailable(RuntimeError):
    """The Excel-compatible formula engine could not be loaded for this workbook."""


def formula_text(value):
    """openpyxl yields str for normal formulas, ArrayFormula objects for CSE formulas."""
    if isinstance(value, str):
        return value if value.startswith("=") else None
    text = getattr(value, "text", None)
    return text if isinstance(text, str) and text.startswith("=") else None


def grid_sheets(wb):
    """Chartsheets have no cells; skip them without crashing on .iter_rows."""
    return [(name, wb[name]) for name in wb.sheetnames if hasattr(wb[name], "iter_rows")]


def is_error(value):
    return isinstance(value, str) and value.startswith(ERROR_VALUES)


def _axis(letter, absolute, index, origin):
    if absolute:
        return f"{letter}{index}"
    offset = index - origin
    return f"{letter}[{offset}]" if offset else letter


def _relative(ref, row, col):
    """A1 reference -> R1C1 relative to the holding cell, so fill-down copies compare equal."""
    sheet, bang, area = ref.rpartition("!")
    parts = area.split(":")
    out = []
    for part in parts:
        cell = CELL_REF.match(part)
        column = COL_REF.match(part) if len(parts) == 2 else None
        line = ROW_REF.match(part) if len(parts) == 2 else None
        if cell:
            out.append(_axis("R", cell.group(3), int(cell.group(4)), row)
                       + _axis("C", cell.group(1), column_index_from_string(cell.group(2).upper()), col))
        elif column:
            out.append(_axis("C", column.group(1), column_index_from_string(column.group(2).upper()), col))
        elif line:
            out.append(_axis("R", line.group(1), int(line.group(2)), row))
        else:
            return ref
    return sheet + bang + ":".join(out)


def analyse(formula, coord):
    """(shape, functions) for one formula.

    The shape rewrites only cell references, so two formulas share a shape exactly
    when one is a fill-down or fill-right copy of the other. Constants, operators and
    string literals stay verbatim: =EXP(0.78*D9) and =EXP(1.05*D9) are different shapes.
    """
    col_letter, row = coordinate_from_string(coord)
    col = column_index_from_string(col_letter)
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return formula, []
    parts, functions = [], []
    for token in tokens:
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            parts.append(_relative(token.value, row, col))
            continue
        if token.type == Token.FUNC and token.subtype == Token.OPEN:
            name = token.value[:-1].upper()
            functions.append(name.replace("_XLFN.", "").replace("_XLWS.", ""))
        parts.append(token.value)
    return "=" + "".join(parts), functions


def _local(element):
    return element.tag.rsplit("}", 1)[-1]


def uncached_formula_cells(path):
    """{sheet: {coords}} of formula cells that hold no calculated result.

    openpyxl reads both a never-calculated formula and a formula whose result is the
    empty string as None. They are different: the second is a real cached value. In
    the file an empty-string result is typed t="str"; an uncalculated formula is not.
    """
    out = {}
    with zipfile.ZipFile(path) as archive:
        rels = {rel.get("Id"): rel.get("Target")
                for rel in ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        for element in ElementTree.fromstring(archive.read("xl/workbook.xml")).iter():
            if _local(element) != "sheet":
                continue
            rel_id = next((v for k, v in element.attrib.items() if k.endswith("}id")), None)
            target = rels.get(rel_id, "")
            part = target.lstrip("/") if target.startswith("/") else "xl/" + target
            if part not in archive.namelist() or "worksheets/" not in part:
                continue
            missing = set()
            with archive.open(part) as handle:
                for _, cell in ElementTree.iterparse(handle):
                    if _local(cell) != "c":
                        continue
                    children = {_local(child): child for child in cell}
                    if "f" in children:
                        value = children.get("v")
                        blank = value is None or not (value.text or "")
                        if blank and cell.get("t") != "str":
                            missing.add(cell.get("r"))
                    cell.clear()
            out[element.get("name")] = missing
    return out


def as_number(value):
    """Excel stores dates and times as serial numbers; openpyxl hands back datetimes."""
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)):
        return to_excel(value)
    return value


def agree(a, b, rtol=1e-9):
    """True when two cell values are the same number, or the same non-number."""
    a, b = as_number(a), as_number(b)
    if isinstance(a, str) or isinstance(b, str):
        return ("" if a is None else a) == ("" if b is None else b)
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return a == b
    if math.isnan(a) or math.isnan(b):
        return False
    return math.isclose(a, b, rel_tol=rtol, abs_tol=1e-12)


def _plain(value):
    """Engine results arrive as 1x1 numpy arrays of numpy scalars, XlError or EMPTY."""
    if hasattr(value, "shape"):
        value = value.ravel()[0] if value.size else None
    if hasattr(value, "item"):
        value = value.item()
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    text = str(value)
    return text if text.startswith("#") else ""


class FormulaEngine:
    """Evaluates the workbook's own formula strings with Excel semantics."""

    def __init__(self, path):
        try:
            import formulas
        except ImportError as exc:
            raise EngineUnavailable("formula engine not installed: pip install formulas") from exc
        try:
            with contextlib.redirect_stderr(io.StringIO()):
                self._model = formulas.ExcelModel().loads(str(path)).finish()
                self._keys = {}
                self._runs = {(): self._read(self._model.calculate())}
        except Exception as exc:
            raise EngineUnavailable(f"formula engine could not load the workbook: {exc}") from exc

    def _read(self, solution):
        out = {}
        for key, cell in solution.items():
            match = ENGINE_KEY.match(key)
            if match:
                ident = (match.group(2), match.group(3) + match.group(4))
                self._keys[ident] = key
                out[ident] = _plain(cell.value)
        return out

    def values(self, inputs=None):
        """{(SHEET, coord): value}, recalculated with `inputs` {(sheet, coord): value} overridden."""
        run = tuple(sorted((sheet.upper(), coord, value) for (sheet, coord), value in (inputs or {}).items()))
        if run not in self._runs:
            overrides = {}
            for sheet, coord, value in run:
                if (sheet, coord) not in self._keys:
                    raise KeyError(f"{sheet}!{coord} is not a cell the engine knows")
                overrides[self._keys[(sheet, coord)]] = value
            with contextlib.redirect_stderr(io.StringIO()):
                self._runs[run] = self._read(self._model.calculate(inputs=overrides))
        return self._runs[run]

    def value(self, sheet, coord, inputs=None):
        return self.values(inputs).get((sheet.upper(), coord))


class FormulaBook:
    """Every formula in a workbook, grouped into distinct shapes, with its cached value."""

    def __init__(self, workbook, cached_from=None):
        """cached_from: a recalculated copy to read cached values from. Formula text
        always comes from `workbook` — recalculating rewrites it (FALSE -> FALSE())."""
        self.path = Path(workbook)
        values_path = Path(cached_from or workbook)
        try:
            wb_f = openpyxl.load_workbook(self.path, data_only=False)
            wb_v = openpyxl.load_workbook(values_path, data_only=True)
        except Exception as exc:
            sys.exit(f"cannot read {self.path}: {exc}")
        uncached = uncached_formula_cells(values_path)
        self.cells = {}
        self.shapes = []
        by_shape = {}
        for name, ws in grid_sheets(wb_f):
            ws_v = wb_v[name]
            missing = uncached.get(name, set())
            for row in ws.iter_rows():
                for cell in row:
                    text = formula_text(cell.value)
                    if text is None:
                        continue
                    coord = cell.coordinate
                    shape, functions = analyse(text, coord)
                    cached = ws_v[coord].value
                    if cached is None and coord not in missing:
                        cached = ""
                    entry = by_shape.get((name, shape))
                    if entry is None:
                        entry = {
                            "id": f"F{len(self.shapes) + 1}", "sheet": name, "cell": coord,
                            "formula": text, "shape": shape, "functions": sorted(set(functions)),
                            "volatile": bool(set(functions) & VOLATILE), "cells": [],
                        }
                        by_shape[(name, shape)] = entry
                        self.shapes.append(entry)
                    entry["cells"].append(coord)
                    self.cells[(name, coord)] = {
                        "formula": text, "shape_id": entry["id"], "cached": cached,
                        "uncached": coord in missing, "array_ref": getattr(cell.value, "ref", None),
                    }
        try:
            self.defined_names = {name: defn.attr_text for name, defn in wb_f.defined_names.items()}
        except Exception:
            self.defined_names = {}
        self._engine = None

    def formula(self, sheet, coord):
        """The formula at a cell exactly as stored, or None for a constant."""
        cell = self.cells.get((sheet, coord))
        return cell["formula"] if cell else None

    def cached(self, sheet, coord):
        cell = self.cells.get((sheet, coord))
        return cell["cached"] if cell else None

    def shape_of(self, sheet, coord):
        cell = self.cells.get((sheet, coord))
        return cell["shape_id"] if cell else None

    @property
    def engine(self):
        if self._engine is None:
            self._engine = FormulaEngine(self.path)
        return self._engine

    def engine_parity(self, rtol=1e-9):
        """Compare the engine's recalculation of every formula with Excel's cached value.

        Returns {"agree": n, "volatile": n, "uncached": n, "disagree": [...]}; each
        disagreement is one distinct shape with its first offending cell.
        """
        values = self.engine.values()
        volatile = {s["id"] for s in self.shapes if s["volatile"]}
        result = {"agree": 0, "volatile": 0, "uncached": 0, "disagree": [], "cells": set()}
        seen = set()
        for (sheet, coord), cell in self.cells.items():
            if cell["shape_id"] in volatile:
                result["volatile"] += 1
            elif cell["uncached"] or is_error(cell["cached"]):
                result["uncached"] += 1
            elif agree(values.get((sheet.upper(), coord)), cell["cached"], rtol):
                result["agree"] += 1
            else:
                result["cells"].add((sheet, coord))
                if cell["shape_id"] not in seen:
                    seen.add(cell["shape_id"])
                    result["disagree"].append({
                        "sheet": sheet, "cell": coord, "formula": cell["formula"],
                        "excel": cell["cached"], "engine": values.get((sheet.upper(), coord)),
                    })
        return result

    def manifest(self):
        """The full formula record that ships beside the converted notebook."""
        return {
            "workbook": self.path.name,
            "sha256": hashlib.sha256(self.path.read_bytes()).hexdigest(),
            "n_formula_cells": len(self.cells),
            "n_distinct_formulas": len(self.shapes),
            "defined_names": self.defined_names,
            "shapes": self.shapes,
            "cells": [{"sheet": sheet, "cell": coord, **cell} for (sheet, coord), cell in self.cells.items()],
        }

    def skeleton(self):
        """A FORMULAS dict to paste into the conversion: one entry per distinct formula."""
        lines = ["FORMULAS = {"]
        for shape in self.shapes:
            span = "" if len(shape["cells"]) == 1 else f"   # {len(shape['cells'])} cells, to {shape['cells'][-1]}"
            lines.append(f"    {shape['sheet'] + '!' + shape['cell']!r}: ({shape['formula']!r},{span}")
            lines.append("        'TODO python expression'),")
        lines.append("}")
        return "\n".join(lines)


def provenance_table(carried, excluded=None):
    """Markdown audit table built from the conversion's FORMULAS map, never typed by hand."""
    def code(text):
        return "`" + str(text).replace("|", "\\|") + "`"

    rows = ["| Excel cell | Formula | Python |", "|---|---|---|"]
    for ref, (excel, python) in carried.items():
        rows.append(f"| {code(ref)} | {code(excel)} | {code(python)} |")
    for ref, reason in (excluded or {}).items():
        rows.append(f"| {code(ref)} | — | **excluded:** {reason} |")
    return "\n".join(rows)


def export_with_formulas(source, out, inputs=None):
    """Write a workbook that still holds every source formula, optionally with new inputs.

    The export carries no cached values; Excel or LibreOffice calculates them on open.
    """
    source = Path(source)
    wb = openpyxl.load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    for (sheet, coord), value in (inputs or {}).items():
        if formula_text(wb[sheet][coord].value):
            raise ValueError(f"{sheet}!{coord} holds a formula; only input cells can be overridden")
        wb[sheet][coord] = value
    wb.save(out)
    before, after = FormulaBook(source), FormulaBook(out)
    lost = [f"{sheet}!{coord}" for (sheet, coord), cell in before.cells.items()
            if after.formula(sheet, coord) != cell["formula"]]
    if lost or len(after.cells) != len(before.cells):
        raise RuntimeError(f"export changed {len(lost)} formulas, e.g. {lost[:5]}")
    return Path(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="command", required=True)
    p_manifest = sub.add_parser("manifest", help="extract every formula to JSON")
    p_manifest.add_argument("workbook")
    p_manifest.add_argument("--out", help="default: <workbook stem>.formulas.json")
    p_manifest.add_argument("--skeleton", action="store_true", help="print a FORMULAS dict to fill in")
    p_engine = sub.add_parser("engine", help="recalculate with the formula engine, compare with cached values")
    p_engine.add_argument("workbook")
    p_export = sub.add_parser("export", help="write a copy with every formula intact")
    p_export.add_argument("workbook")
    p_export.add_argument("out")
    args = ap.parse_args()

    if args.command == "export":
        print(f"wrote {export_with_formulas(args.workbook, args.out)} with every formula intact")
        return

    book = FormulaBook(args.workbook)
    if args.command == "manifest":
        out = Path(args.out) if args.out else book.path.with_suffix(".formulas.json")
        out.write_text(json.dumps(book.manifest(), indent=1, default=str, ensure_ascii=False), encoding="utf-8")
        print(f"{len(book.cells)} formula cells, {len(book.shapes)} distinct formulas -> {out}")
        if args.skeleton:
            print(book.skeleton())
        return

    try:
        parity = book.engine_parity()
    except EngineUnavailable as exc:
        sys.exit(str(exc))
    print(f"engine reproduces {parity['agree']} cached formula results; "
          f"{parity['volatile']} volatile and {parity['uncached']} uncached or broken cells not compared")
    for d in parity["disagree"]:
        print(f"  DISAGREES  {d['sheet']}!{d['cell']}  {d['formula'][:70]}")
        print(f"             engine={d['engine']!r}  excel={d['excel']!r}")


if __name__ == "__main__":
    main()
