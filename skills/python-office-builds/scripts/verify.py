#!/usr/bin/env python3
"""verify.py — post-build QA for .docx / .xlsx. Prints a JSON report.

  python verify.py out.docx              # lint + render pages → out_pages/page-N.png
  python verify.py out.xlsx              # lint + LibreOffice recalc (in place) + error scan
  python verify.py out.xlsx --render     # also render sheets to PNG
  python verify.py out.xlsx --no-recalc  # lint only (file untouched)

Exit code: 0 clean, 1 problems found, 2 tool failure. Never treat "rendered"
as "correct" — open the PNGs and look.
"""
from __future__ import annotations

import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

BUILTIN_OFFICE = Path(os.environ.get("OFFICEKIT_BUILTIN", "/mnt/skills/public/xlsx/scripts"))  # claude.ai sandbox helpers
ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")
BANNED = re.compile(r"\b(XLOOKUP|XMATCH|SORT|SORTBY|FILTER|UNIQUE|SEQUENCE|LET|LAMBDA|RANDARRAY)\s*\(", re.I)
NEEDS_XLFN = re.compile(r"(?<!_xlfn\.)\b(TEXTJOIN|CONCAT|IFS|SWITCH|MAXIFS|MINIFS)\s*\(", re.I)
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
Sub UpdateIndexesAndSave()
  Dim idx, i
  idx = ThisComponent.getDocumentIndexes()
  For i = 0 To idx.getCount() - 1
    idx.getByIndex(i).update()
  Next i
  ThisComponent.getTextFields().refresh()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""


# ───────────────────────────── soffice plumbing ─────────────────────────────
def soffice_bin() -> str | None:
    for cand in (os.environ.get("SOFFICE"), shutil.which("soffice"), shutil.which("libreoffice"),
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                 r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if cand and Path(cand).exists():
            return cand
    return None


def soffice_env() -> dict:
    if (BUILTIN_OFFICE / "office" / "soffice.py").exists():
        sys.path.insert(0, str(BUILTIN_OFFICE))
        try:
            from office.soffice import get_soffice_env  # handles sandbox socket shim
            return get_soffice_env()
        except Exception:
            pass
    env = os.environ.copy()
    if platform.system() == "Linux":
        env.setdefault("SAL_USE_VCLPLUGIN", "svp")
    return env


def run_soffice(args: list[str], profile: Path, timeout: int = 120) -> subprocess.CompletedProcess:
    exe = soffice_bin()
    if not exe:
        raise FileNotFoundError("LibreOffice (soffice) not found — set $SOFFICE or install it "
                                "(macOS: brew install --cask libreoffice)")
    cmd = [exe, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore", *args]
    return subprocess.run(cmd, capture_output=True, text=True, env=soffice_env(), timeout=timeout)


def to_pdf(path: Path, outdir: Path) -> Path:
    with tempfile.TemporaryDirectory(prefix="lo_") as prof:
        r = run_soffice(["--convert-to", "pdf", "--outdir", str(outdir), str(path)], Path(prof))
    pdf = outdir / (path.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError(f"PDF conversion failed: {r.stderr.strip() or r.stdout.strip()}")
    return pdf


def pdf_to_png(pdf: Path, outdir: Path, dpi: int = 80, max_pages: int = 12) -> list[str]:
    outdir.mkdir(parents=True, exist_ok=True)
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-l", str(max_pages), str(pdf),
                        str(outdir / "page")], check=True, capture_output=True)
    else:
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(pdf)
            for i, page in enumerate(doc):
                if i >= max_pages:
                    break
                page.get_pixmap(dpi=dpi).save(outdir / f"page-{i + 1}.png")
        except ImportError:
            return []
    return sorted(str(p) for p in outdir.glob("page-*.png"))


def pdf_pages(pdf: Path) -> int | None:
    if shutil.which("pdfinfo"):
        out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        m = re.search(r"Pages:\s+(\d+)", out)
        return int(m.group(1)) if m else None
    return None


# ─────────────────────────────────── DOCX ───────────────────────────────────
def verify_docx(path: Path, render: bool = True) -> dict:
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(path)
    issues, info = [], {}
    paras = doc.paragraphs
    info["paragraphs"] = len(paras)
    info["tables"] = len(doc.tables)
    info["images"] = len(doc.inline_shapes)
    heads = [p.text for p in paras if p.style.name.startswith("Heading")]
    info["headings"] = heads[:40]
    sec = doc.sections[0]
    info["page_in"] = [round(sec.page_width / 914400, 2), round(sec.page_height / 914400, 2)]

    run = blanks = 0
    for p in paras:
        run = run + 1 if not p.text.strip() and not p._p.xpath(".//w:drawing|.//w:br") else 0
        blanks = max(blanks, run)
        if "\n" in p.text:
            issues.append(f"literal newline inside paragraph: {p.text[:50]!r} — use separate paragraphs")
        if re.match(r"^\s*[•\-\*]\s", p.text) and not p.style.name.startswith("List"):
            issues.append(f"fake bullet character: {p.text[:50]!r} — use 'List Bullet' style")
        if re.search(r"\{\{.*?\}\}|\bTODO\b|\bTBD\b|lorem ipsum|\[INSERT", p.text, re.I):
            issues.append(f"placeholder text left in: {p.text[:60]!r}")
    if blanks >= 3:
        issues.append(f"{blanks} consecutive empty paragraphs — use spacing/page breaks instead")
    for ti, t in enumerate(doc.tables):
        for row in t.rows:
            for c in row.cells:
                if c._tc.tcPr is None or c._tc.tcPr.find(qn("w:tcW")) is None:
                    issues.append(f"table {ti + 1}: cell without explicit width (Google Docs/LibreOffice may collapse it)")
                    break
            else:
                continue
            break
        for c in t.rows[0].cells if t.rows else []:
            shd = c._tc.tcPr.find(qn("w:shd")) if c._tc.tcPr is not None else None
            if shd is not None and shd.get(qn("w:val")) == "solid":
                issues.append(f"table {ti + 1}: w:shd val=solid renders black — use clear")
                break
    rep = {"file": str(path), "type": "docx", "info": info, "issues": issues}
    if render:  # NB: `issues` list is shared with rep, appends below still land
        outdir = path.with_name(path.stem + "_pages")
        shutil.rmtree(outdir, ignore_errors=True)
        outdir.mkdir()
        src = path
        if doc.element.body.xpath(".//w:instrText[contains(., 'TOC')]"):
            src = outdir / path.name  # preview copy; the deliverable keeps Word's updateFields
            shutil.copy(path, src)
            res = run_macro(src, "UpdateIndexesAndSave")
            if "error" in res:
                rep_note = "TOC refresh failed; preview shows an empty TOC (Word fills it on open)"
                issues.append(rep_note)
        pdf = to_pdf(src, outdir)
        if src != path:
            src.unlink()
        rep["pdf"] = str(pdf)
        rep["pages"] = pdf_pages(pdf)
        rep["png"] = pdf_to_png(pdf, outdir)
    return rep


# ─────────────────────────────────── XLSX ───────────────────────────────────
def run_macro(path: Path, sub: str, timeout: int = 90) -> dict:
    """Open `path` in headless LibreOffice, run Module1.<sub>, which stores in place."""
    with tempfile.TemporaryDirectory(prefix="lo_") as prof_s:
        prof = Path(prof_s)
        run_soffice(["--terminate_after_init"], prof, timeout=timeout)
        mdir = prof / "user" / "basic" / "Standard"
        if not mdir.exists():
            return {"error": "LibreOffice profile not created; macro not run"}
        (mdir / "Module1.xba").write_text(MACRO)
        before = path.stat().st_mtime_ns
        r = run_soffice([f"vnd.sun.star.script:Standard.Module1.{sub}?language=Basic&location=application",
                         str(path.resolve())], prof, timeout=timeout)
        if path.stat().st_mtime_ns == before:
            return {"error": "LibreOffice did not rewrite the file — close other LibreOffice "
                             f"instances and retry. stderr: {r.stderr.strip()[:300]}"}
    return {"ok": True}


def recalc(path: Path, timeout: int = 90) -> dict:
    res = run_macro(path, "RecalculateAndSave", timeout)
    return res if "error" in res else {"recalculated": True}


def verify_xlsx(path: Path, do_recalc: bool = True, render: bool = False) -> dict:
    from openpyxl import load_workbook
    issues, info = [], {}
    wbf = load_workbook(path)  # formulas
    if any(n.startswith("xl/externalLinks/") for n in __import__("zipfile").ZipFile(path).namelist()):
        issues.append("workbook has external links — recalc would destroy them; skipping recalc")
        do_recalc = False
    n_formulas = 0
    for ws in wbf.worksheets:
        info.setdefault("sheets", []).append(f"{ws.title} [{ws.dimensions}]")
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                v = getattr(v, "text", v)
                if isinstance(v, str) and v.startswith("="):
                    n_formulas += 1
                    if BANNED.search(v):
                        issues.append(f"{ws.title}!{c.coordinate}: dynamic-array/LET function won't evaluate outside Excel 365: {v[:60]}")
                    if NEEDS_XLFN.search(v):
                        issues.append(f"{ws.title}!{c.coordinate}: needs _xlfn. prefix: {v[:60]}")
                    if re.search(r"(?<![!'\w])[A-Za-z][\w ]* [\w ]*!\$?[A-Z]+\$?\d", v) and "'" not in v:
                        issues.append(f"{ws.title}!{c.coordinate}: sheet name with space must be quoted: {v[:60]}")
    info["formulas"] = n_formulas
    rep = {"file": str(path), "type": "xlsx", "info": info, "issues": issues}

    if n_formulas and do_recalc:
        res = None
        if (BUILTIN_OFFICE / "recalc.py").exists():  # prefer the sandbox's hardened recalc
            p = subprocess.run([sys.executable, str(BUILTIN_OFFICE / "recalc.py"), str(path.resolve()), "90"],
                               capture_output=True, text=True, cwd=str(BUILTIN_OFFICE))
            try:
                res = json.loads(p.stdout)
            except json.JSONDecodeError:
                res = None
        if res is None:
            res = recalc(path)
        rep["recalc"] = {k: res.get(k) for k in ("status", "total_formulas", "total_errors", "error", "recalculated") if k in res}
        if "error" in res:
            issues.append("RECALC FAILED: " + str(res["error"]))
    if n_formulas and not do_recalc:
        rep["warning"] = "formulas have no cached values until opened in Excel or recalculated"

    wbv = load_workbook(path, data_only=True)
    errs, empties = [], 0
    for ws in wbv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERRORS:
                    errs.append(f"{ws.title}!{c.coordinate}={c.value}")
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and wbv[ws.title][c.coordinate].value is None:
                    empties += 1
    if errs:
        issues.append(f"{len(errs)} formula error cell(s): " + ", ".join(errs[:25]))
    if do_recalc and n_formulas and empties and "error" not in rep.get("recalc", {}):
        rep["note"] = f"{empties} formula cell(s) evaluate to blank — confirm that's intended"
    if render:
        outdir = path.with_name(path.stem + "_pages")
        shutil.rmtree(outdir, ignore_errors=True)
        outdir.mkdir()
        tmp = outdir / path.name
        shutil.copy(path, tmp)
        pdf = to_pdf(tmp, outdir)
        tmp.unlink()
        rep["png"] = pdf_to_png(pdf, outdir)
    return rep


def main() -> int:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    if not args:
        print(__doc__)
        return 2
    path = Path(args[0]).resolve()
    t0 = time.time()
    try:
        if path.suffix.lower() == ".docx":
            rep = verify_docx(path, render="--no-render" not in flags)
        elif path.suffix.lower() in (".xlsx", ".xlsm"):
            rep = verify_xlsx(path, do_recalc="--no-recalc" not in flags, render="--render" in flags)
        else:
            print(json.dumps({"error": f"unsupported: {path.suffix}"}))
            return 2
    except Exception as e:  # tool failure, not a content problem
        print(json.dumps({"file": str(path), "error": f"{type(e).__name__}: {e}"}, indent=2))
        return 2
    rep["seconds"] = round(time.time() - t0, 1)
    rep["status"] = "clean" if not rep["issues"] else "issues_found"
    print(json.dumps(rep, indent=2, ensure_ascii=False))
    return 0 if not rep["issues"] else 1


if __name__ == "__main__":
    sys.exit(main())
